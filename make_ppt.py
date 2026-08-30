from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUTPUT = "Pause-for-Paws-student-presentation.pptx"
BG = RGBColor(243, 240, 232)
INK = RGBColor(23, 34, 30)
GREEN = RGBColor(36, 92, 67)
MINT = RGBColor(199, 228, 195)
YELLOW = RGBColor(240, 184, 75)
CORAL = RGBColor(223, 110, 81)
MUTED = RGBColor(104, 116, 109)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

slides = [
    ("The problem", ["Animals cross roads to find food, water, and shelter.", "Vehicle collisions can hurt animals and people.", "Some roads have more wildlife crossings than others.", "Our goal is to help drivers slow down in those places."], "Pause for Paws uses past crossing information to identify roads where drivers should pay extra attention."),
    ("Where the information comes from", ["A coordinate tells us where an event happened.", "An event can include a date, animal type, and confidence score.", "Our current data is synthetic practice data, not real Waze data.", "Real data needs permission from an agency or approved partner."], "We never pretend that practice data is real. That is important for safety and honesty."),
    ("How the map works", ["The map places historical events near their coordinates.", "Nearby events are grouped into a longer road corridor.", "Repeated events create a higher historical risk score.", "The map shows patterns, not a guaranteed live animal location."], "A corridor pattern is more useful than saying an animal is definitely at one exact spot right now."),
    ("The Pause for Paws message", ["A vehicle approaches a known risk area.", "The system can show: Pause for Paws.", "The message says to slow down and watch both shoulders.", "Historical data cannot prove an animal is there right now."], "This is a safety reminder, not a promise that an animal is currently on the road."),
    ("What makes our idea useful", ["It combines locations with animal and seasonal information.", "It turns many old reports into easy-to-understand patterns.", "Agencies can use patterns to plan signs, fencing, or crossings.", "A simulator lets us test before using real driver alerts."], "Our idea helps people make better decisions while keeping the message simple."),
    ("Our next steps", ["Use verified data with permission.", "Test with a wildlife agency and transportation department.", "Check accuracy and make sure alerts do not distract drivers.", "Protect API keys and private location information.", "Get legal review before public use."], "Ranger questions to prepare for: Is the data real? How accurate are the coordinates? Could the message distract drivers? Does it change routes? How will you prove it helps?"),
]


def add_text(slide, text, left, top, width, height, size, color=INK, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_base(slide, number, title):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = BG
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(.23), prs.slide_height)
    accent.fill.solid(); accent.fill.fore_color.rgb = CORAL; accent.line.fill.background()
    add_text(slide, "PAUSE FOR PAWS", .65, .45, 3, .25, 10, GREEN, True, "Aptos Mono")
    add_text(slide, f"0{number}", 11.75, .45, .8, .25, 10, MUTED, False, "Aptos Mono", PP_ALIGN.RIGHT)
    add_text(slide, title, .65, 1.15, 8.7, .75, 34, INK, True, "Aptos Display")
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, .65, 2.05, 1.0, .06)
    line.fill.solid(); line.fill.fore_color.rgb = GREEN; line.line.fill.background()


def add_bullets(slide, bullets):
    box = slide.shapes.add_textbox(Inches(.9), Inches(2.5), Inches(8.2), Inches(3.9))
    frame = box.text_frame; frame.clear(); frame.word_wrap = True
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.space_after = Pt(17)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(22)
        paragraph.font.color.rgb = INK
        paragraph._p.get_or_add_pPr().insert(0, paragraph._p.get_or_add_pPr()._new_buChar()) if False else None
        paragraph.text = "•  " + item
    return box

for number, (title, bullets, speaker_note) in enumerate(slides, 1):
    slide = prs.slides.add_slide(blank)
    add_base(slide, number, title)
    add_bullets(slide, bullets)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.45), Inches(2.45), Inches(3.05), Inches(2.95))
    card.fill.solid(); card.fill.fore_color.rgb = GREEN; card.line.fill.background()
    add_text(slide, "OUR SIGNAL", 9.78, 2.8, 2.2, .25, 10, MINT, True, "Aptos Mono")
    signal = "Slow down.\nSomething is moving." if number == 1 else "Pause for\nPaws"
    add_text(slide, signal, 9.78, 3.35, 2.3, 1.25, 27, RGBColor(255, 255, 255), True, "Aptos Display")
    add_text(slide, "Historical pattern • not a live tracker", 9.78, 5.0, 2.35, .3, 10, MINT, False, "Aptos Mono")
    try:
        slide.notes_slide.notes_text_frame.text = speaker_note
    except AttributeError:
        pass

prs.save(OUTPUT)
print(f"Created {OUTPUT}")
